"""Tests for /api/notify-interest (public POST + admin GET + admin export)."""
import os
import pytest
import requests
from openpyxl import load_workbook
from io import BytesIO

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', 'https://sacred-rudraksha-hub.preview.emergentagent.com').rstrip('/')
ADMIN_EMAIL = "admin@rudralife.com"
ADMIN_PASSWORD = "Rudralife@2026"


@pytest.fixture(scope="module")
def api_client():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


@pytest.fixture(scope="module")
def admin_token(api_client):
    r = api_client.post(f"{BASE_URL}/api/admin/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD})
    assert r.status_code == 200, r.text
    return r.json()["token"]


@pytest.fixture(scope="module")
def auth_headers(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


VALID_NOTIFY = {
    "full_name": "TEST_Notify_Lead",
    "email": "test_notify@example.com",
    "dial_code": "+91",
    "phone": "9876543210",
    "interested_city": "Chennai",
    "exhibition_type": "domestic"
}


class TestNotifyInterestPublic:
    def test_create_notify_interest_success(self, api_client):
        r = api_client.post(f"{BASE_URL}/api/notify-interest", json=VALID_NOTIFY)
        assert r.status_code == 200, r.text
        data = r.json()
        assert "id" in data and isinstance(data["id"], str) and len(data["id"]) > 0
        assert data["full_name"] == VALID_NOTIFY["full_name"]
        assert data["email"] == VALID_NOTIFY["email"]
        assert data["interested_city"] == "Chennai"
        assert data["exhibition_type"] == "domestic"
        assert "_id" not in data
        pytest.notify_id = data["id"]

    def test_create_notify_interest_international(self, api_client):
        payload = dict(VALID_NOTIFY)
        payload["email"] = "test_notify_intl@example.com"
        payload["interested_city"] = "Dubai"
        payload["exhibition_type"] = "international"
        r = api_client.post(f"{BASE_URL}/api/notify-interest", json=payload)
        assert r.status_code == 200
        data = r.json()
        assert data["exhibition_type"] == "international"
        assert data["interested_city"] == "Dubai"

    def test_invalid_email_returns_422(self, api_client):
        bad = dict(VALID_NOTIFY)
        bad["email"] = "not-an-email"
        r = api_client.post(f"{BASE_URL}/api/notify-interest", json=bad)
        assert r.status_code == 422

    def test_missing_required_field_returns_422(self, api_client):
        bad = dict(VALID_NOTIFY)
        bad.pop("interested_city")
        r = api_client.post(f"{BASE_URL}/api/notify-interest", json=bad)
        assert r.status_code == 422


class TestNotifyInterestAdmin:
    def test_list_requires_auth(self, api_client):
        r = api_client.get(f"{BASE_URL}/api/admin/notify-interest")
        assert r.status_code == 401

    def test_list_returns_items(self, api_client, auth_headers):
        r = api_client.get(f"{BASE_URL}/api/admin/notify-interest", headers=auth_headers)
        assert r.status_code == 200
        data = r.json()
        assert "total" in data and "items" in data
        assert isinstance(data["items"], list)
        assert data["total"] >= 1
        # Verify our created TEST_ leads are present
        emails = [i.get("email") for i in data["items"]]
        assert "test_notify@example.com" in emails
        # No mongo _id leaking
        for item in data["items"]:
            assert "_id" not in item

    def test_filter_by_city(self, api_client, auth_headers):
        r = api_client.get(f"{BASE_URL}/api/admin/notify-interest?city=Chennai", headers=auth_headers)
        assert r.status_code == 200
        data = r.json()
        for item in data["items"]:
            assert item["interested_city"].lower() == "chennai"

    def test_filter_by_type(self, api_client, auth_headers):
        r = api_client.get(f"{BASE_URL}/api/admin/notify-interest?exhibition_type=international", headers=auth_headers)
        assert r.status_code == 200
        data = r.json()
        for item in data["items"]:
            assert item["exhibition_type"] == "international"

    def test_filter_by_far_future_date_returns_empty(self, api_client, auth_headers):
        r = api_client.get(f"{BASE_URL}/api/admin/notify-interest?date_from=2099-01-01", headers=auth_headers)
        assert r.status_code == 200
        assert r.json()["total"] == 0

    def test_export_xlsx_requires_auth(self, api_client):
        r = api_client.get(f"{BASE_URL}/api/admin/notify-interest/export")
        assert r.status_code == 401

    def test_export_xlsx_success(self, api_client, auth_headers):
        r = api_client.get(f"{BASE_URL}/api/admin/notify-interest/export", headers=auth_headers)
        assert r.status_code == 200
        assert r.headers.get("content-type", "").startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        cd = r.headers.get("content-disposition", "")
        assert "rudralife_notifyme_" in cd and ".xlsx" in cd
        # Validate xlsx
        wb = load_workbook(BytesIO(r.content))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) >= 2  # header + at least one data row
        headers = rows[0]
        assert headers[0] == "Created At"
        assert "Interested City" in headers
        assert "Exhibition Type" in headers
