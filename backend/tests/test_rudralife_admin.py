"""Backend API tests for Rudralife admin panel (auth + exhibitions CRUD + leads + export)."""
import os
import io
import pytest
import requests

BASE_URL = os.environ.get(
    "REACT_APP_BACKEND_URL",
    "https://sacred-rudraksha-hub.preview.emergentagent.com",
).rstrip("/")

ADMIN_EMAIL = "admin@rudralife.com"
ADMIN_PASSWORD = "Rudralife@2026"


@pytest.fixture(scope="module")
def client():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


@pytest.fixture(scope="module")
def token(client):
    r = client.post(f"{BASE_URL}/api/admin/login",
                    json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD})
    assert r.status_code == 200, f"Login failed: {r.status_code} {r.text}"
    data = r.json()
    assert "token" in data and isinstance(data["token"], str) and len(data["token"]) > 20
    assert data["admin"]["email"] == ADMIN_EMAIL
    return data["token"]


@pytest.fixture(scope="module")
def auth_client(client, token):
    s = requests.Session()
    s.headers.update({
        "Content-Type": "application/json",
        "Authorization": f"Bearer {token}",
    })
    return s


# ─────────── Public exhibitions seed ───────────
class TestPublicExhibitions:
    def test_public_list_returns_seeded(self, client):
        r = client.get(f"{BASE_URL}/api/exhibitions")
        assert r.status_code == 200
        data = r.json()
        assert isinstance(data, list)
        assert len(data) >= 20, f"Expected ≥20 seeded, got {len(data)}"
        domestic = [e for e in data if e.get("type") == "domestic"]
        intl = [e for e in data if e.get("type") == "international"]
        assert len(domestic) >= 16, f"Expected ≥16 domestic, got {len(domestic)}"
        assert len(intl) >= 4, f"Expected ≥4 international, got {len(intl)}"
        # No _id leakage
        for e in data[:5]:
            assert "_id" not in e
            assert "id" in e and "name" in e and "status" in e


# ─────────── Admin auth ───────────
class TestAdminAuth:
    def test_login_wrong_password(self, client):
        r = client.post(f"{BASE_URL}/api/admin/login",
                        json={"email": ADMIN_EMAIL, "password": "WRONG_PASS"})
        assert r.status_code == 401

    def test_login_unknown_email(self, client):
        r = client.post(f"{BASE_URL}/api/admin/login",
                        json={"email": "nobody@rudralife.com", "password": ADMIN_PASSWORD})
        assert r.status_code == 401

    def test_login_correct(self, token):
        assert isinstance(token, str) and len(token) > 20

    def test_me_with_token(self, auth_client):
        r = auth_client.get(f"{BASE_URL}/api/admin/me")
        assert r.status_code == 200
        data = r.json()
        assert data["email"] == ADMIN_EMAIL
        assert "password_hash" not in data
        assert "_id" not in data

    def test_me_without_token(self, client):
        r = client.get(f"{BASE_URL}/api/admin/me")
        assert r.status_code == 401


# ─────────── Admin exhibitions CRUD ───────────
class TestAdminExhibitions:
    def test_list_requires_auth(self, client):
        r = client.get(f"{BASE_URL}/api/admin/exhibitions")
        assert r.status_code == 401

    def test_list_with_auth(self, auth_client):
        r = auth_client.get(f"{BASE_URL}/api/admin/exhibitions")
        assert r.status_code == 200
        assert isinstance(r.json(), list)
        assert len(r.json()) >= 20

    def test_full_lifecycle(self, auth_client, client):
        # CREATE
        payload = {
            "name": "TEST_City_Lifecycle",
            "type": "domestic",
            "image": "https://example.com/x.jpg",
            "status": "soon",
            "start_date": "2026-06-01",
            "end_date": "2026-06-05",
            "timings": "10:00 AM to 8:00 PM",
            "venue": "Test Hotel",
            "address": "Test Address, City",
            "order": 999,
        }
        r = auth_client.post(f"{BASE_URL}/api/admin/exhibitions", json=payload)
        assert r.status_code == 200, r.text
        created = r.json()
        eid = created["id"]
        assert created["name"] == payload["name"]
        assert created["status"] == "soon"

        # Public should reflect new exhibition
        pub = client.get(f"{BASE_URL}/api/exhibitions").json()
        assert any(e["id"] == eid for e in pub), "New exhibition not in public list"

        # PUT update
        upd = dict(payload)
        upd["name"] = "TEST_City_Updated"
        upd["venue"] = "Updated Hotel"
        r = auth_client.put(f"{BASE_URL}/api/admin/exhibitions/{eid}", json=upd)
        assert r.status_code == 200
        assert r.json()["name"] == "TEST_City_Updated"
        assert r.json()["venue"] == "Updated Hotel"

        # PATCH status -> live
        r = auth_client.patch(f"{BASE_URL}/api/admin/exhibitions/{eid}/status",
                              json={"status": "live"})
        assert r.status_code == 200
        assert r.json()["status"] == "live"

        # PATCH invalid value
        r = auth_client.patch(f"{BASE_URL}/api/admin/exhibitions/{eid}/status",
                              json={"status": "bogus"})
        assert r.status_code == 400

        # GET public reflects status change
        pub = client.get(f"{BASE_URL}/api/exhibitions").json()
        match = next((e for e in pub if e["id"] == eid), None)
        assert match and match["status"] == "live"

        # DELETE
        r = auth_client.delete(f"{BASE_URL}/api/admin/exhibitions/{eid}")
        assert r.status_code == 200
        assert r.json().get("deleted") is True

        # Confirm 404 on second delete
        r = auth_client.delete(f"{BASE_URL}/api/admin/exhibitions/{eid}")
        assert r.status_code == 404

        # Public list no longer contains it
        pub = client.get(f"{BASE_URL}/api/exhibitions").json()
        assert not any(e["id"] == eid for e in pub)

    def test_create_requires_auth(self, client):
        r = client.post(f"{BASE_URL}/api/admin/exhibitions",
                        json={"name": "x", "type": "domestic"})
        assert r.status_code == 401


# ─────────── Admin registrations (read-only + filters + export) ───────────
class TestAdminRegistrations:
    def _seed(self, client, exhibition_id, exhibition_city, visit_date, marker):
        payload = {
            "full_name": f"TEST_Lead_{marker}",
            "email": f"test_lead_{marker}@example.com",
            "dial_code": "+91",
            "phone": "9000000000",
            "city": "Mumbai",
            "country": "India",
            "exhibition_id": exhibition_id,
            "exhibition_city": exhibition_city,
            "visit_date": visit_date,
            "message": f"marker_{marker}",
        }
        r = client.post(f"{BASE_URL}/api/register", json=payload)
        assert r.status_code == 200, r.text
        return r.json()

    def test_requires_auth(self, client):
        r = client.get(f"{BASE_URL}/api/admin/registrations")
        assert r.status_code == 401

    def test_list_and_filters(self, auth_client, client):
        # Need a real exhibition_id
        exs = client.get(f"{BASE_URL}/api/exhibitions").json()
        first = exs[0]
        # Seed two registrations with different visit dates
        r1 = self._seed(client, first["id"], first["name"], "2026-05-10", "fA")
        r2 = self._seed(client, first["id"], first["name"], "2026-05-11", "fB")

        # List all
        r = auth_client.get(f"{BASE_URL}/api/admin/registrations")
        assert r.status_code == 200
        body = r.json()
        assert "total" in body and "items" in body
        assert body["total"] >= 2
        assert isinstance(body["items"], list)
        for it in body["items"][:5]:
            assert "_id" not in it

        # Filter by exhibition_id
        r = auth_client.get(
            f"{BASE_URL}/api/admin/registrations?exhibition_id={first['id']}"
        )
        assert r.status_code == 200
        for it in r.json()["items"]:
            assert it["exhibition_id"] == first["id"]

        # Filter by visit_date exact match
        r = auth_client.get(
            f"{BASE_URL}/api/admin/registrations?visit_date=2026-05-10"
        )
        assert r.status_code == 200
        items = r.json()["items"]
        assert len(items) >= 1
        for it in items:
            assert it["visit_date"] == "2026-05-10"

        # Filter by date_from in far future returns 0
        r = auth_client.get(
            f"{BASE_URL}/api/admin/registrations?date_from=2099-01-01"
        )
        assert r.status_code == 200
        assert r.json()["total"] == 0

    def test_no_delete_endpoint(self, auth_client, client):
        # Confirm DELETE on /api/admin/registrations/<id> is not allowed
        exs = client.get(f"{BASE_URL}/api/exhibitions").json()
        reg = self._seed(client, exs[0]["id"], exs[0]["name"], "2026-05-12", "noDel")
        r = auth_client.delete(f"{BASE_URL}/api/admin/registrations/{reg['id']}")
        assert r.status_code in (404, 405), \
            f"Registrations should NEVER be deletable; got {r.status_code}"

    def test_export_excel(self, auth_client):
        r = auth_client.get(f"{BASE_URL}/api/admin/registrations/export")
        assert r.status_code == 200
        ct = r.headers.get("content-type", "")
        assert "spreadsheetml.sheet" in ct, f"Bad Content-Type: {ct}"
        cd = r.headers.get("content-disposition", "")
        assert "attachment" in cd and ".xlsx" in cd, f"Bad CD: {cd}"
        # Parse the workbook to ensure it's valid
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        # Header row exists
        headers = [c.value for c in ws[1]]
        assert "Full Name" in headers and "Email" in headers
        assert ws.max_row >= 2  # at least one data row from seeded test data

    def test_export_requires_auth(self, client):
        r = client.get(f"{BASE_URL}/api/admin/registrations/export")
        assert r.status_code == 401
